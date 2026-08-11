from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "manuais" / "manual_operacao_servidor_detec.docx"

BLUE = "145DA0"
DARK_BLUE = "0A2B5E"
LIGHT_BLUE = "EAF3FB"
PALE_BLUE = "F5F9FD"
GRAY = "5E6B7A"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D7DEE7"
GREEN = "147D64"
PALE_GREEN = "EAF7F2"
ORANGE = "A85A00"
PALE_ORANGE = "FFF4E5"
RED = "B42318"
PALE_RED = "FDECEC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    # left/right keeps the document compatible with older desktop Word builds.
    for margin, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(section, text: str) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("DeTec Server  |  Manual operacional  |  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    keep_with_next(paragraph)
    return paragraph


def add_body(doc: Document, text: str, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)



def add_steps(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(item)


def add_code(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(code.strip().splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string("1F2937")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, title: str, body: str, tone: str = "info") -> None:
    palette = {
        "info": (LIGHT_BLUE, BLUE),
        "success": (PALE_GREEN, GREEN),
        "warning": (PALE_ORANGE, ORANGE),
        "danger": (PALE_RED, RED),
    }
    fill, accent = palette[tone]
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(title + "\n")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(accent)
    body_run = paragraph.add_run(body)
    body_run.font.color.rgb = RGBColor.from_string("26364A")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if row_index % 2:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for idx, text in enumerate(values):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.paragraphs[0].add_run(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_checklist(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        run = paragraph.add_run("☐  " + item)
        run.font.name = "Segoe UI Symbol"


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("1C2735")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3)

    heading_specs = {
        "Title": (28, DARK_BLUE, 0, 16),
        "Subtitle": (13, GRAY, 0, 10),
        "Heading 1": (17, DARK_BLUE, 18, 8),
        "Heading 2": (13, BLUE, 13, 6),
        "Heading 3": (11, DARK_BLUE, 9, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    section = doc.sections[0]
    set_repeat_header(section, "DEIGO TECNOLOGIA  •  OPERAÇÃO DO SERVIDOR LOCAL")
    add_page_number(section)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(1.05)
    run = p.add_run("DEIGO TECNOLOGIA")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph(style="Title")
    title.add_run("Guia operacional\nDeTec Server")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Instalação, configuração, manutenção e recuperação do servidor local")

    doc.add_paragraph().paragraph_format.space_after = Inches(0.5)
    add_callout(
        doc,
        "Referência de produção",
        "Edição 2.0 • Atualizado em 11/08/2026 • DeTec Server 0.1.11 • DeTec PDV 0.1.10",
        "info",
    )
    add_body(doc, "Público: técnico de implantação, suporte autorizado e administrador master.")
    add_body(doc, "Ambiente principal: Windows, PostgreSQL e serviço DeigoVarejoServidorLocal.")
    doc.add_paragraph().paragraph_format.space_after = Inches(1.0)
    add_callout(
        doc,
        "Regra de ouro",
        "Antes de alterar configuração, banco, serviço ou versão, faça backup e registre o que será mudado. Nunca edite o arquivo .env que está dentro do ZIP de distribuição para corrigir somente um cliente.",
        "warning",
    )
    add_page_break(doc)

    add_heading(doc, "Como usar este guia", 1)
    add_body(doc, "Siga primeiro o fluxo rápido. As seções seguintes detalham cada operação e os procedimentos de recuperação.")
    add_table(
        doc,
        ["Necessidade", "Seção"],
        [
            ["Instalar ou atualizar o servidor", "3. Instalação e atualização"],
            ["Alterar uma configuração", "5. Edição segura do .env"],
            ["Servidor não abre", "7. Diagnóstico e incidentes"],
            ["Criar administrador master", "6. Usuários administrativos"],
            ["Fazer ou restaurar backup", "9 e 10"],
            ["Publicar ou reconfigurar PDV", "11. Aplicativos e terminais"],
            ["Remover a instalação", "13. Desinstalação segura"],
        ],
        [3600, 5760],
    )
    add_heading(doc, "Fluxo rápido", 2)
    add_steps(doc, [
        "Abra o PowerShell como administrador.",
        "Confirme o serviço e o healthcheck.",
        "Faça backup antes de qualquer manutenção.",
        "Aplique a alteração pelo instalador aprovado ou no arquivo instalado.",
        "Reinicie o serviço e valide login, versão, arquivos estáticos e acesso pela rede.",
        "Registre data, técnico, motivo e resultado.",
    ])
    add_heading(doc, "Referência rápida", 2)
    add_table(
        doc,
        ["Item", "Valor padrão"],
        [
            ["Código instalado", r"C:\DeTecServer\app"],
            ["Configuração do cliente", r"C:\DeTecServer\app\.env"],
            ["Dados, mídia e estáticos", r"C:\ProgramData\DeigoVarejo\Dados"],
            ["Backups", r"C:\ProgramData\DeigoVarejo\Backups"],
            ["Logs do serviço", r"C:\ProgramData\DeigoVarejo\ServidorLocal\logs"],
            ["Serviço Windows", "DeigoVarejoServidorLocal"],
            ["Acesso local de exemplo", "http://127.0.0.1:8001/login/"],
        ],
        [3300, 6060],
    )
    add_page_break(doc)

    add_heading(doc, "1. Segurança antes da manutenção", 1)
    add_callout(doc, "PowerShell elevado", "Serviço, firewall, tarefa agendada e instalação exigem PowerShell aberto com Executar como administrador.", "warning")
    add_checklist(doc, [
        "Confirmar qual empresa, servidor, IP, porta e versão serão alterados.",
        "Confirmar que o backup mais recente existe e possui arquivo .sha256.",
        "Guardar uma cópia externa do backup antes de atualização importante.",
        "Não compartilhar .env, senhas, certificados, chaves de ativação ou backup em repositório público.",
        "Não abrir a porta 5432 do PostgreSQL para a rede; somente o ERP deve acessar o banco.",
        "Planejar janela de manutenção e avisar os caixas antes de reiniciar o serviço.",
    ])
    add_heading(doc, "O que nunca fazer", 2)
    add_bullets(doc, [
        "Apagar C:\\ProgramData\\DeigoVarejo antes de validar um backup.",
        "Editar arquivos dentro do ZIP e esperar que a instalação já existente seja atualizada.",
        "Executar manage.py runserver como servidor permanente da loja.",
        "Trocar SECRET_KEY ou FISCAL_CERTIFICATE_KEY sem plano de migração e cópia segura.",
        "Copiar o .env de uma empresa para outra.",
    ])

    add_heading(doc, "2. Serviço e acesso diário", 1)
    add_heading(doc, "Consultar, iniciar, parar e reiniciar", 2)
    add_code(doc, "Get-Service DeigoVarejoServidorLocal\nStart-Service DeigoVarejoServidorLocal\nStop-Service DeigoVarejoServidorLocal\nRestart-Service DeigoVarejoServidorLocal")
    add_body(doc, "Após iniciar ou reiniciar, valide o endereço configurado para a instalação:")
    add_code(doc, "Invoke-WebRequest http://127.0.0.1:8001/login/ -UseBasicParsing")
    add_body(doc, "Um status HTTP 200 indica que a tela de login respondeu. Redirecionamentos 301/302 também podem ser normais em rotas autenticadas.")
    add_heading(doc, "Descobrir o endereço da rede", 2)
    add_code(doc, "ipconfig\nGet-NetIPAddress -AddressFamily IPv4 | Where-Object {$_.IPAddress -notlike '169.254*'}")
    add_callout(doc, "IP fixo", "Reserve o IP no roteador ou configure endereço estático. Se o IP mudar, PDVs e acessos administrativos podem deixar de encontrar o servidor.", "info")
    add_page_break(doc)

    add_heading(doc, "3. Instalação e atualização", 1)
    add_heading(doc, "Instalação recomendada no Windows", 2)
    add_steps(doc, [
        "Copie o ZIP oficial para uma pasta local e confira o SHA-256 publicado.",
        "Extraia todo o ZIP. Não execute o instalador de dentro do compactador.",
        "Abra o PowerShell como administrador na pasta extraída.",
        "Execute Instalar DeTec Server.exe. O script PowerShell permanece disponível para suporte.",
        "Informe IP, porta e senhas solicitadas. A primeira instalação pode preparar PostgreSQL e o runtime Python portátil.",
        "Ao final, confira serviço, healthcheck, login, arquivos estáticos e acesso em outro computador da rede.",
    ])
    add_heading(doc, "Atualização de uma instalação existente", 2)
    add_body(doc, "Use um pacote de versão superior. O instalador detecta a instalação existente, atualiza o código e preserva dados e configuração local. Ainda assim, faça backup antes.")
    add_checklist(doc, [
        "Confirmar o número da versão do pacote e seu SHA-256.",
        "Executar backup e guardar uma cópia fora do servidor.",
        "Executar o instalador como administrador no mesmo IP e porta.",
        "Não usar -Force sem orientação técnica; ele é reservado para reinstalação controlada.",
        "Validar migrations, collectstatic, serviço, login, Central do PDV e Central do Admin.",
    ])
    add_callout(doc, "GitHub não é o instalador do cliente", "Em produção, entregue o pacote aprovado. O repositório é usado para desenvolvimento e construção; não faça git clone no servidor da loja como procedimento padrão.", "warning")

    add_heading(doc, "4. Backup antes de mudar", 1)
    add_body(doc, "Execute a partir da pasta instalada:")
    add_code(doc, "Set-Location C:\\DeTecServer\\app\n.\\scripts\\backup_local.ps1 -ValidarSomente\n.\\scripts\\backup_local.ps1 -IncluirLogs")
    add_body(doc, "O backup padrão vai para C:\\ProgramData\\DeigoVarejo\\Backups e gera arquivo de integridade .sha256. A opção -ValidarSomente confere fontes e ferramentas sem criar o backup.")
    add_heading(doc, "Agendar backup diário", 2)
    add_code(doc, ".\\scripts\\register_backup_task.ps1 -Horario 02:30 -RetencaoDias 15 -Force")
    add_callout(doc, "Cópia externa", "Um backup no mesmo disco não protege contra falha física, ransomware ou perda do servidor. Mantenha uma cópia criptografada em destino externo controlado.", "danger")
    add_page_break(doc)

    add_heading(doc, "5. Edição segura do .env", 1)
    add_callout(doc, "Arquivo correto", "Edite C:\\DeTecServer\\app\\.env na instalação do cliente. Não edite o .env.example nem um arquivo dentro do ZIP para corrigir somente aquele servidor.", "info")
    add_heading(doc, "Procedimento", 2)
    add_steps(doc, [
        "Abra o PowerShell como administrador.",
        "Crie uma cópia do .env com data e hora.",
        "Abra o arquivo instalado no Bloco de Notas.",
        "Altere apenas as chaves necessárias, salve em UTF-8 e feche o editor.",
        "Reinicie o serviço.",
        "Confirme status, healthcheck, login e a função relacionada à alteração.",
    ])
    add_code(doc, "$stamp = Get-Date -Format 'yyyyMMdd-HHmmss'\nCopy-Item C:\\DeTecServer\\app\\.env \"C:\\DeTecServer\\app\\.env.backup-$stamp\"\nnotepad C:\\DeTecServer\\app\\.env\nRestart-Service DeigoVarejoServidorLocal\nGet-Service DeigoVarejoServidorLocal\nInvoke-WebRequest http://127.0.0.1:8001/login/ -UseBasicParsing")
    add_heading(doc, "Regras do arquivo", 2)
    add_bullets(doc, [
        "Use uma única linha KEY=value por chave e não crie chaves duplicadas.",
        "Não inclua espaços antes do nome da chave; preserve o formato do valor existente.",
        "Não publique o .env no GitHub, e-mail ou chamado sem remover os segredos.",
        "Não troque SECRET_KEY, senhas do banco ou chaves fiscais por tentativa e erro.",
        "Uma alteração no .env só entra em vigor depois que o processo do serviço reinicia.",
    ])
    add_heading(doc, "Versões e publicação dos aplicativos", 2)
    add_table(
        doc,
        ["Chave", "Uso"],
        [
            ["LOCAL_SERVER_VERSION", "Versão do servidor instalada."],
            ["PDV_DESKTOP_VERSION", "Versão publicada do DeTec PDV."],
            ["PDV_DESKTOP_MIN_VERSION", "Menor versão aceita nos terminais."],
            ["PDV_DESKTOP_REQUIRE_SIGNED_INSTALLER", "Exige artefato assinado quando true."],
        ],
        [3600, 5760],
    )
    add_callout(doc, "Padrão para novas lojas", "Se a mudança deve valer para todas as futuras instalações, altere a fonte do projeto/instalador, rode os testes e gere um novo pacote. Alterar somente o .env instalado afeta apenas aquele cliente.", "warning")
    add_heading(doc, "Voltar atrás", 2)
    add_code(doc, "Stop-Service DeigoVarejoServidorLocal\nCopy-Item C:\\DeTecServer\\app\\.env.backup-AAAAMMDD-HHMMSS C:\\DeTecServer\\app\\.env -Force\nStart-Service DeigoVarejoServidorLocal")
    add_page_break(doc)

    add_heading(doc, "6. Usuários administrativos", 1)
    add_heading(doc, "Criar o primeiro superusuário", 2)
    add_code(doc, "Set-Location C:\\DeTecServer\\app\n.\\.venv\\Scripts\\python.exe manage.py createsuperuser")
    add_body(doc, "Informe usuário, e-mail e senha quando solicitado. Use uma conta individual; não compartilhe o login master entre operadores.")
    add_heading(doc, "Alterar senha", 2)
    add_code(doc, ".\\.venv\\Scripts\\python.exe manage.py changepassword NOME_DO_USUARIO")
    add_heading(doc, "Listar superusuários", 2)
    add_code(doc, ".\\.venv\\Scripts\\python.exe manage.py shell -c \"from django.contrib.auth import get_user_model; print(list(get_user_model().objects.filter(is_superuser=True).values_list('username', flat=True)))\"")
    add_callout(doc, "Separação de papéis", "O superadmin da Deigo Tecnologia administra licenças e implantação. O administrador da empresa enxerga somente sua empresa e filiais; o operador de caixa não deve acessar telas administrativas.", "info")

    add_heading(doc, "7. Diagnóstico e incidentes", 1)
    add_heading(doc, "Diagnóstico padrão", 2)
    add_code(doc, "Set-Location C:\\DeTecServer\\app\n.\\scripts\\test_local_server_service.ps1 -HealthUrl http://127.0.0.1:8001/login/")
    add_heading(doc, "Consultar logs", 2)
    add_code(doc, "Get-ChildItem C:\\ProgramData\\DeigoVarejo\\ServidorLocal\\logs\nGet-Content C:\\ProgramData\\DeigoVarejo\\ServidorLocal\\logs\\*.log -Tail 120")
    add_table(
        doc,
        ["Sintoma", "Ação inicial"],
        [
            ["Serviço parado", "Consultar logs; depois Start-Service."],
            ["Healthcheck falha", "Confirmar porta do .env/serviço, log e conflito de porta."],
            ["Tela sem estilo", "Executar collectstatic, conferir STATIC_ROOT e reiniciar."],
            ["Erro 403 CSRF", "Conferir URL, cookies, relógio, ALLOWED_HOSTS e CSRF_TRUSTED_ORIGINS."],
            ["Outro computador não acessa", "Conferir IP fixo, perfil de rede privada, firewall e porta do ERP."],
            ["Versão do PDV diverge", "Conferir .env, manifesto do artefato e reiniciar o serviço."],
            ["Banco indisponível", "Conferir serviço PostgreSQL e credenciais sem expor a porta 5432."],
        ],
        [3000, 6360],
    )
    add_page_break(doc)

    add_heading(doc, "8. Manutenção Django", 1)
    add_body(doc, "Execute sempre com o Python do ambiente virtual instalado:")
    add_code(doc, "Set-Location C:\\DeTecServer\\app\n.\\.venv\\Scripts\\python.exe manage.py check\n.\\.venv\\Scripts\\python.exe manage.py showmigrations\n.\\.venv\\Scripts\\python.exe manage.py migrate --noinput\n.\\.venv\\Scripts\\python.exe manage.py collectstatic --noinput")
    add_callout(doc, "Servidor permanente", "O serviço Windows usa Waitress. manage.py runserver é apenas para desenvolvimento e não deve substituir o serviço da loja.", "warning")
    add_heading(doc, "Teste manual temporário", 2)
    add_code(doc, ".\\scripts\\run_local_server.ps1 -Bind 127.0.0.1 -Port 8010 -NoMigrate")
    add_body(doc, "Use uma porta livre e encerre com Ctrl+C. Não deixe a janela como solução definitiva.")

    add_heading(doc, "9. Backup e retenção", 1)
    add_table(
        doc,
        ["Operação", "Comando"],
        [
            ["Validar fontes", r".\scripts\backup_local.ps1 -ValidarSomente"],
            ["Backup completo", r".\scripts\backup_local.ps1 -IncluirLogs"],
            ["Destino alternativo", r".\scripts\backup_local.ps1 -Destino D:\Backups -RetencaoDias 30"],
            ["Agendar diariamente", r".\scripts\register_backup_task.ps1 -Horario 02:30 -RetencaoDias 15 -Force"],
        ],
        [3000, 6360],
    )
    add_checklist(doc, [
        "O ZIP ou AES foi criado e possui .sha256.",
        "O manifesto identifica banco, mídia e data.",
        "A cópia externa foi atualizada.",
        "A tarefa agendada está ativa e o último resultado foi bem-sucedido.",
        "Um restore de teste foi executado periodicamente.",
    ])

    add_heading(doc, "10. Restauração", 1)
    add_callout(doc, "Operação destrutiva", "A restauração altera dados em uso. Pare os caixas, faça backup do estado atual e valide o arquivo antes de confirmar.", "danger")
    add_heading(doc, "Validar sem restaurar", 2)
    add_code(doc, ".\\scripts\\restore_local_backup.ps1 -BackupPath \"D:\\Backups\\supermercado-local-AAAAMMDD-HHMMSS.zip\" -Port 8001 -ValidarSomente")
    add_heading(doc, "Restaurar", 2)
    add_code(doc, ".\\scripts\\restore_local_backup.ps1 -BackupPath \"D:\\Backups\\supermercado-local-AAAAMMDD-HHMMSS.zip\" -Port 8001 -ConfirmarRestauracao")
    add_body(doc, "Para backup criptografado, informe BACKUP_ENCRYPTION_PASSPHRASE no processo autorizado. Nunca registre a senha em documento público ou histórico de comando compartilhado.")
    add_checklist(doc, [
        "SHA-256 validado.",
        "Backup do estado atual concluído.",
        "Loja em janela de manutenção.",
        "Serviço voltou a responder.",
        "Login, empresas, estoque, vendas, mídia e relatórios conferidos.",
    ])
    add_page_break(doc)

    add_heading(doc, "11. Aplicativos e terminais", 1)
    add_heading(doc, "DeTec PDV", 2)
    add_bullets(doc, [
        "O artefato deve ser construído, publicado e compatível com PDV_DESKTOP_VERSION.",
        "O admin master cadastra e licencia o terminal; a chave de ativação não deve ser exposta em logs ou chamados.",
        "Se a chave for rotacionada, a anterior é invalidada e a ação deve ficar auditada.",
        "No computador do caixa, a configuração fica em %LOCALAPPDATA%\\DeTecPDV\\config.json.",
        "Para reconfigurar o servidor ou uma nova chave, use a opção Nova chave/Configurar do aplicativo; não edite segredos manualmente.",
    ])
    add_heading(doc, "DeTec Admin", 2)
    add_bullets(doc, [
        "É a casca administrativa do mesmo ERP e não substitui o servidor.",
        "A configuração local fica em %LOCALAPPDATA%\\DeTecAdmin\\config.json.",
        "O login continua sendo validado pelo ERP; o aplicativo não concede permissões extras.",
    ])
    add_heading(doc, "Publicação após atualização", 2)
    add_steps(doc, [
        "Gere os artefatos DeTecPDV e DeTecAdmin na máquina de build.",
        "Publique-os no pacote/servidor com manifestos e SHA-256 corretos.",
        "Confira as versões no .env e reinicie o serviço.",
        "Abra as Centrais de download e confirme que o botão está liberado.",
        "Homologue em um terminal piloto antes de distribuir para todos os caixas.",
    ])

    add_heading(doc, "12. Atualização e rollback", 1)
    add_heading(doc, "Após cada atualização", 2)
    add_checklist(doc, [
        "Serviço Running e healthcheck saudável.",
        "manage.py check sem erros e migrations aplicadas.",
        "CSS, imagens e arquivos enviados pelo usuário aparecem.",
        "Versão do servidor e versões dos aplicativos coincidem com os manifestos.",
        "Login de superadmin e administrador da empresa funcionam.",
        "Um PDV piloto abre, autentica e consulta produtos.",
        "Backup pós-atualização criado.",
    ])
    add_heading(doc, "Rollback", 2)
    add_body(doc, "Se a atualização falhar, preserve logs, interrompa novas operações, restaure o .env anterior e use o pacote/backup homologado. Não misture código de uma versão com migrations ou artefatos de outra sem roteiro técnico.")
    add_callout(doc, "Sem apagar evidências", "Copie os logs e registre o erro antes de reinstalar. Isso permite corrigir a causa em vez de apenas ocultar o sintoma.", "info")
    add_page_break(doc)

    add_heading(doc, "13. Desinstalação segura", 1)
    add_callout(doc, "Não basta apagar a pasta", "O sistema instala serviço, regra de firewall, banco, tarefa de backup, código e dados. Remova cada componente de forma controlada.", "danger")
    add_steps(doc, [
        "Faça e valide um backup; copie-o para outro equipamento.",
        "Pare os caixas e registre a autorização para desinstalar.",
        "Abra PowerShell como administrador e remova o serviço pelo script.",
        "Remova tarefas agendadas e regras de firewall identificadas para o DeTec Server.",
        "Remova o código em C:\\DeTecServer somente após confirmar o backup.",
        "Remova C:\\ProgramData\\DeigoVarejo somente se os dados e backups não precisarem ser preservados.",
        "Trate PostgreSQL separadamente; não desinstale se ele for compartilhado por outra aplicação.",
    ])
    add_code(doc, "Set-Location C:\\DeTecServer\\app\n.\\scripts\\uninstall_local_server_service.ps1 -RemoveServiceFiles")
    add_body(doc, "A opção -RemoveServiceFiles remove apenas os arquivos do wrapper em ProgramData\\DeigoVarejo\\ServidorLocal. Ela não apaga automaticamente banco, mídia e backups.")
    add_heading(doc, "Aplicativos dos computadores", 2)
    add_body(doc, "Desinstale DeTec PDV e DeTec Admin em Aplicativos instalados do Windows. Se necessário e autorizado, remova também as configurações locais em %LOCALAPPDATA%\\DeTecPDV e %LOCALAPPDATA%\\DeTecAdmin.")

    add_heading(doc, "14. Checklist de entrega", 1)
    add_checklist(doc, [
        "Pacote, manifesto e SHA-256 conferidos.",
        "IP fixo, porta, firewall privado e acesso pela rede testados.",
        "PostgreSQL acessível somente pelo servidor.",
        "Serviço configurado para iniciar automaticamente e identidade dedicada confirmada.",
        ".env protegido, sem chaves duplicadas e com cópia segura externa dos segredos.",
        "Superadmin criado; perfis da empresa e operadores testados.",
        "Backup automático registrado e restore de teste concluído.",
        "Central do PDV e Central do Admin publicam artefatos corretos.",
        "Terminal piloto ativado e fluxo básico homologado.",
        "Fiscal, TEF, impressora, balança e gaveta permanecem em homologação até teste real.",
        "Manual, versão instalada, data e responsável registrados.",
    ])
    add_heading(doc, "Registro da manutenção", 2)
    add_table(
        doc,
        ["Campo", "Preenchimento"],
        [
            ["Empresa / filial", ""],
            ["Servidor / IP / porta", ""],
            ["Versão anterior / nova", ""],
            ["Backup / SHA-256", ""],
            ["Alteração realizada", ""],
            ["Testes executados", ""],
            ["Técnico / data / aceite", ""],
        ],
        [3000, 6360],
    )

    add_heading(doc, "Apêndice A. Comandos essenciais", 1)
    add_code(doc, "# Serviço\nGet-Service DeigoVarejoServidorLocal\nRestart-Service DeigoVarejoServidorLocal\n\n# Healthcheck\nInvoke-WebRequest http://127.0.0.1:8001/login/ -UseBasicParsing\n\n# Projeto\nSet-Location C:\\DeTecServer\\app\n.\\.venv\\Scripts\\python.exe manage.py check\n.\\.venv\\Scripts\\python.exe manage.py migrate --noinput\n.\\.venv\\Scripts\\python.exe manage.py collectstatic --noinput\n\n# Usuários\n.\\.venv\\Scripts\\python.exe manage.py createsuperuser\n.\\.venv\\Scripts\\python.exe manage.py changepassword USUARIO\n\n# Backup\n.\\scripts\\backup_local.ps1 -ValidarSomente\n.\\scripts\\backup_local.ps1 -IncluirLogs\n\n# Diagnóstico\n.\\scripts\\test_local_server_service.ps1 -HealthUrl http://127.0.0.1:8001/login/")
    add_callout(doc, "Suporte", "Ao abrir um chamado, informe empresa, versão, horário do erro, operação executada e anexe logs sem senhas, certificados, chaves ou dados pessoais desnecessários.", "success")

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.core_properties.title = "Guia operacional DeTec Server"
    document.core_properties.subject = "Instalação, configuração, manutenção e recuperação do servidor local"
    document.core_properties.author = "Deigo Tecnologia"
    document.core_properties.keywords = "DeTec Server, servidor local, PostgreSQL, backup, PDV, operação"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
