"""Gera o manual operacional de cadastro de produtos."""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "manuais" / "manual_cadastro_produtos.docx"
SCREEN = ROOT / "docs" / "manuais" / "imagens" / "tela_cadastro_produto_classificacao.png"
NAVY, BLUE, PALE, WHITE = "082B69", "0B5BE7", "F5F8FC", "FFFFFF"
TEXT, MUTED, BORDER = "17233B", "5D6B82", "CAD6E8"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def margins(cell, value=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    elem = OxmlElement("w:tblHeader")
    elem.set(qn("w:val"), "true")
    tr_pr.append(elem)


def title(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    return p


def callout(doc, label, body, fill="EAF2FF", accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    margins(cell, 170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    p.add_run(body)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def steps(doc, items):
    for number, heading, body in items:
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        left, right = table.rows[0].cells
        left.width, right.width = Cm(1.1), Cm(15.7)
        shade(left, BLUE)
        shade(right, PALE)
        for cell in (left, right):
            margins(cell, 140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(number))
        r.bold, r.font.size = True, Pt(14)
        r.font.color.rgb = RGBColor.from_string(WHITE)
        p = right.paragraphs[0]
        r = p.add_run(heading)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        p.add_run(f"\n{body}")
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def field_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(4.1), Cm(8.0), Cm(4.7))
    for idx, (cell, label) in enumerate(zip(table.rows[0].cells, ("Campo", "Como preencher", "Exemplo"))):
        cell.width = widths[idx]
        shade(cell, NAVY)
        margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
    repeat_header(table.rows[0])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, (cell, value) in enumerate(zip(cells, row)):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            if row_idx % 2:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            if idx == 0:
                r.bold = True
    return table


def checklist(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        mark = p.add_run("☐ ")
        mark.bold = True
        mark.font.color.rgb = RGBColor.from_string(BLUE)
        p.add_run(item)


def setup(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 31, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 14, BLUE)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("DEIGO TECNOLOGIA  |  DEIGO VAREJO")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Manual operacional • Cadastro de produtos")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup(doc)

    band = doc.add_table(rows=1, cols=1)
    cell = band.cell(0, 0)
    shade(cell, NAVY)
    margins(cell, 360)
    r = cell.paragraphs[0].add_run("DEIGO VAREJO")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(42)
    p.add_run("Manual de cadastro\nde produtos")
    p = doc.add_paragraph()
    r = p.add_run("Do primeiro código de barras à liberação para PDV e marketplace")
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph("Versão 1.0 • 31 de julho de 2026")
    doc.add_paragraph().paragraph_format.space_after = Pt(75)
    callout(doc, "Objetivo", "Cadastrar o produto uma única vez, com identificação, categoria, embalagem, preço, imagem e dados fiscais consistentes. Os PDVs do servidor local usam a alteração imediatamente; instalações híbridas recebem atualização automática.")
    doc.add_page_break()

    title(doc, "1. Visão geral da tela")
    doc.add_paragraph("A tela Novo produto concentra as informações operacionais. O mockup destaca a classificação comercial, que é a categoria organizada em níveis.")
    doc.add_picture(str(SCREEN), width=Cm(17.5))
    p = doc.add_paragraph("Figura 1 — Mockup didático da tela de cadastro.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    callout(doc, "Categoria = classificação comercial", "Selecione sempre o nível mais específico. O sistema guarda o caminho completo, por exemplo: Alimentos > Mercearia > Massas > Macarrão.")
    doc.add_page_break()

    title(doc, "2. Antes de cadastrar")
    steps(doc, [
        (1, "Separe os dados", "Tenha EAN, descrição, unidade, custo, preço, NCM, origem e tributação."),
        (2, "Confira a classificação", "Cadastre departamento, seção, grupo ou subgrupo que ainda não existir."),
        (3, "Defina a unidade base", "Escolha como o saldo será controlado: UN, KG, G, L, M, CX, FD ou PCT."),
        (4, "Valide a parte fiscal", "Confirme os dados com a contabilidade. Não preencha tributação por suposição."),
    ])
    title(doc, "Como acessar", 2)
    doc.add_paragraph("Abra Cadastros > Produtos e selecione Novo produto. O usuário precisa ter permissão de cadastro.")
    title(doc, "Como criar a categoria", 2)
    field_table(doc, [
        ("Departamento", "Maior divisão comercial.", "Alimentos"),
        ("Seção", "Área dentro do departamento.", "Mercearia"),
        ("Grupo", "Família de itens semelhantes.", "Massas"),
        ("Subgrupo", "Classificação mais específica.", "Macarrão"),
    ])
    callout(doc, "Regra", "A ordem é Departamento > Seção > Grupo > Subgrupo. Não pule níveis ao escolher a classificação superior.", "FFF5DE", "9A5B00")
    doc.add_page_break()

    title(doc, "3. Identificação do produto")
    field_table(doc, [
        ("Nome", "Descrição clara, com marca, conteúdo ou peso quando ajudar.", "Arroz Branco Tipo 1 5 kg"),
        ("Código de barras", "EAN/GTIN principal da unidade base. Deve ser único.", "7891234567890"),
        ("Código interno", "Informe ou deixe vazio para geração automática.", "PRD-000125"),
        ("Tipo comercial", "Mercadoria, insumo, serviço, imobilizado ou material de consumo.", "Mercadoria para revenda"),
        ("Classificação", "Selecione a categoria mais específica disponível.", "Alimentos > Mercearia > Cereais > Arroz"),
        ("Marca", "Selecione a marca comercial quando houver.", "Marca Exemplo"),
        ("Descrição", "Detalhes adicionais e conteúdo para marketplace.", "Pacote de 5 kg"),
    ])
    callout(doc, "Atenção", "Código interno e código de barras são identificadores diferentes. Ambos podem ser usados nas buscas.", "FFF5DE", "9A5B00")
    doc.add_page_break()

    title(doc, "4. Unidades, embalagens e códigos")
    doc.add_paragraph("A unidade base define o saldo. A unidade de compra e o fator de conversão informam como a embalagem recebida se transforma nesse saldo.")
    field_table(doc, [
        ("Unidade base", "Unidade usada no estoque e na venda.", "UN para lata; KG para banana"),
        ("Unidade de compra", "Forma normalmente recebida do fornecedor.", "CX"),
        ("Quantidade na base", "Número de unidades-base dentro da embalagem.", "12 para caixa com 12 latas"),
        ("Peso líquido/bruto", "Informe em kg quando disponível.", "5,000 kg / 5,120 kg"),
        ("EAN adicional", "Código de pacote, fardo ou caixa com seu fator.", "EAN da caixa = 12 UN"),
    ])
    title(doc, "Exemplo: caixa de leite", 2)
    steps(doc, [
        (1, "Unidade base", "Use UN, pois cada caixa de leite é vendida individualmente."),
        (2, "Unidade de compra", "Use CX quando o fornecedor entrega caixa fechada."),
        (3, "Fator", "Informe 12 se a caixa de compra contém 12 unidades."),
        (4, "EAN da caixa", "Cadastre o EAN adicional com fator 12. Ao bipar esse EAN, entram 12 unidades."),
    ])
    callout(doc, "Produto pesável", "Para banana, carne ou outro item por peso, use KG como unidade base e marque Produto pesável.")
    callout(doc, "PLU e setor de balança", "Para itens pesáveis, selecione o setor da empresa e informe um PLU único. Cadastre também tara e validade quando a etiquetadora utilizar esses dados.", "EAF2FF", "155EEF")
    doc.add_page_break()

    title(doc, "5. Preços, estoque, venda e imagens")
    field_table(doc, [
        ("Preço de custo", "Custo de referência. Entradas alimentam o custo médio do estoque.", "R$ 20,00"),
        ("Margem desejada", "Percentual sobre o preço de venda usado para calcular uma sugestão.", "30,00%"),
        ("Preço sugerido", "Referência calculada; não substitui automaticamente o preço praticado.", "R$ 28,57"),
        ("Preço de venda", "Preço normal cobrado no PDV.", "R$ 29,90"),
        ("Preço promocional", "Use somente quando houver promoção aplicável.", "R$ 26,90"),
        ("Estoque mínimo", "Ponto para alerta e sugestão de reposição.", "10 UN"),
        ("Vendido no PDV", "Marque para liberar nos caixas.", "Sim"),
        ("Marketplace", "Marque para disponibilizar na venda online.", "Sim"),
        ("Imagem principal", "PNG, JPG ou JPEG usado como capa.", "produto.jpg"),
        ("Galeria", "Fotos adicionais com legenda e ordem.", "Frente, verso e tabela"),
    ])
    callout(doc, "Fornecedores vinculados", "Registre fornecedor principal, código usado por ele e último custo cotado. A lista é isolada por empresa.", "EAF2FF", "155EEF")
    callout(doc, "Preço antigo no estoque", "Nova compra atualiza o custo médio, mas não muda automaticamente o preço de venda. O reajuste deve ser decidido pela gestão.")
    callout(doc, "Marketplace", "Use foto nítida, fundo neutro e produto inteiro no enquadramento.", "EAF8EF", "148345")
    doc.add_page_break()

    title(doc, "6. Dados fiscais")
    doc.add_paragraph("Os campos fiscais sustentam a emissão de NFC-e e devem refletir a mercadoria e o regime tributário.")
    field_table(doc, [
        ("NCM", "Código fiscal com 8 dígitos. Obrigatório para NFC-e.", "10063021"),
        ("CEST", "Código com 7 dígitos quando aplicável.", "1704900"),
        ("Origem", "Origem conforme tabela fiscal.", "0 - Nacional"),
        ("CST ICMS", "Use quando aplicável ao regime da empresa.", "00"),
        ("CSOSN", "Use quando aplicável ao Simples Nacional.", "102"),
        ("Alíquota ICMS", "Percentual definido pela regra fiscal.", "18,00%"),
    ])
    callout(doc, "Obrigatório", "Confirme NCM, CEST, CST/CSOSN, origem e alíquota com a contabilidade. Categoria comercial não substitui classificação fiscal.", "FFF0F0", "B42318")
    title(doc, "7. Salvar e sincronizar")
    steps(doc, [
        (1, "Revise", "Confira códigos, unidade, conversão, preços, imagem e fiscal."),
        (2, "Salve", "O sistema valida duplicidades e formatos obrigatórios."),
        (3, "Teste a busca", "Procure por nome, código interno e todos os EANs."),
        (4, "Confira o PDV", "No servidor local é imediato; no híbrido, o snapshot é automático e sem carga manual."),
    ])
    doc.add_page_break()

    title(doc, "8. Checklist antes de liberar")
    checklist(doc, [
        "Nome identifica produto e apresentação.",
        "EAN principal não pertence a outro produto.",
        "Código interno foi informado ou será gerado.",
        "Tipo comercial está correto.",
        "Classificação usa o nível mais específico.",
        "Unidade base corresponde ao estoque real.",
        "Unidade de compra e conversão foram conferidas.",
        "EANs adicionais representam corretamente as embalagens.",
        "Preços e estoque mínimo foram revisados.",
        "Produto pesável foi marcado quando necessário.",
        "PLU, setor, tara e validade foram conferidos quando o produto usa balança.",
        "Fornecedor principal e último custo foram revisados."
        "Imagem está em PNG, JPG ou JPEG.",
        "Liberação para PDV e marketplace está correta.",
        "Dados fiscais foram validados pela contabilidade.",
    ])
    title(doc, "Erros comuns", 2)
    field_table(doc, [
        ("EAN duplicado", "Código já usado em outro produto.", "Pesquise antes de cadastrar."),
        ("Conversão incorreta", "Caixa entra como uma unidade.", "Use fator 12 para 12 UN."),
        ("Categoria genérica", "Produto fica em nível muito amplo.", "Use subgrupo quando existir."),
        ("Fiscal incompleto", "Produto não fica pronto para NFC-e.", "Confirme NCM e tributação."),
        ("Imagem quebrada", "Arquivo incompatível ou inválido.", "Use PNG, JPG ou JPEG."),
    ])
    callout(doc, "Suporte", "Ao reportar erro, informe código interno, EAN, filial, usuário, horário e mensagem exibida.")
    doc.core_properties.title = "Manual de cadastro de produtos"
    doc.core_properties.author = "Deigo Tecnologia"
    doc.core_properties.subject = "Procedimento operacional do Deigo Varejo"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
